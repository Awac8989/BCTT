import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('BAO_CAO_THUC_TAP_100_TRANG_DOAN_MINH_QUAN.docx')
for i in range(min(25, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    print(f"P[{i}] (text len={len(p.text.strip())}): {repr(p.text[:60])}")
