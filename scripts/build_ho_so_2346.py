# -*- coding: utf-8 -*-
"""
Script build trọn bộ Hồ sơ công vụ: Xử lý Đơn phản ánh số 2346/UBND ngày 15/10/2026 của bà Nguyễn Thị Mộng Đẹp
Chuẩn thể thức và kỹ thuật trình bày văn bản hành chính theo Nghị định số 30/2020/NĐ-CP của Chính phủ.
"""

import os
import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_font(run, font_name="Times New Roman", size_pt=13, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

def set_para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=3, line_spacing=1.15, first_indent=0, left_indent=0):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_indent:
        pf.first_line_indent = Cm(first_indent)
    else:
        pf.first_line_indent = Cm(0)
    if left_indent:
        pf.left_indent = Cm(left_indent)
    else:
        pf.left_indent = Cm(0)

def set_table_borders_none(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def set_cell_margins(cell, top=0, bottom=40, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_page_setup(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)

def add_header_table(doc, org_line1, org_line2, doc_num, date_str, is_cong_van=False, trich_yeu_cv=""):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table)
    
    col_w_left = Cm(7.2)
    col_w_right = Cm(9.3)
    
    row = table.rows[0]
    cell_left = row.cells[0]
    cell_right = row.cells[1]
    cell_left.width = col_w_left
    cell_right.width = col_w_right
    set_cell_margins(cell_left, top=0, bottom=20, left=0, right=10)
    set_cell_margins(cell_right, top=0, bottom=20, left=10, right=0)
    
    # Left column: Cơ quan ban hành & Số ký hiệu
    p_org = cell_left.paragraphs[0]
    set_para_format(p_org, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=1, line_spacing=1.0)
    if org_line1:
        r1 = p_org.add_run(org_line1 + "\n")
        set_font(r1, size_pt=12, bold=False)
    if org_line2:
        r2 = p_org.add_run(org_line2)
        set_font(r2, size_pt=12, bold=True)
        
        # Đường kẻ ngang dưới tên cơ quan
        p_line_org = cell_left.add_paragraph()
        set_para_format(p_line_org, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, line_spacing=1.0)
        r_line1 = p_line_org.add_run("──────────")
        set_font(r_line1, size_pt=9, bold=True)
    
    # Số ký hiệu
    if doc_num:
        p_num = cell_left.add_paragraph()
        set_para_format(p_num, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1, line_spacing=1.0)
        r_num = p_num.add_run(doc_num)
        set_font(r_num, size_pt=12, bold=False)
        
    # Nếu là công văn có trích yếu dưới số
    if is_cong_van and trich_yeu_cv:
        p_ty = cell_left.add_paragraph()
        set_para_format(p_ty, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=0, line_spacing=1.0)
        r_ty = p_ty.add_run(trich_yeu_cv)
        set_font(r_ty, size_pt=11, bold=False, italic=True)

    # Right column: Quốc hiệu & Tiêu ngữ & Ngày tháng
    p_qh = cell_right.paragraphs[0]
    set_para_format(p_qh, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=1, line_spacing=1.0)
    rqh1 = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    set_font(rqh1, size_pt=12, bold=True)
    rqh2 = p_qh.add_run("Độc lập – Tự do – Hạnh phúc")
    set_font(rqh2, size_pt=13, bold=True)
    
    p_line_qh = cell_right.add_paragraph()
    set_para_format(p_line_qh, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, line_spacing=1.0)
    r_line2 = p_line_qh.add_run("─────────────────")
    set_font(r_line2, size_pt=9, bold=True)
    
    if date_str:
        p_date = cell_right.add_paragraph()
        set_para_format(p_date, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=0, line_spacing=1.0)
        r_date = p_date.add_run(date_str)
        set_font(r_date, size_pt=13, italic=True)

    # Spacer
    p_spacer = doc.add_paragraph()
    set_para_format(p_spacer, space_before=0, space_after=4, line_spacing=1.0)
    return table

def add_doc_title(doc, title, subtitle=""):
    p_t = doc.add_paragraph()
    set_para_format(p_t, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=2, line_spacing=1.15)
    r_t = p_t.add_run(title)
    set_font(r_t, size_pt=14, bold=True)
    
    if subtitle:
        p_st = doc.add_paragraph()
        set_para_format(p_st, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=10, line_spacing=1.15)
        r_st = p_st.add_run(subtitle)
        set_font(r_st, size_pt=12, italic=True, bold=True)

def add_kinh_gui(doc, lines):
    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=4, space_after=6, line_spacing=1.15, first_indent=0)
    r_kg = p.add_run("Kính gửi: ")
    set_font(r_kg, size_pt=13, bold=True)
    
    if len(lines) == 1:
        r_txt = p.add_run(lines[0])
        set_font(r_txt, size_pt=13, bold=False)
    else:
        for idx, line in enumerate(lines):
            if idx > 0:
                p_sub = doc.add_paragraph()
                set_para_format(p_sub, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=2, space_after=2, line_spacing=1.15, first_indent=0, left_indent=1.8)
                r_dash = p_sub.add_run("– ")
                set_font(r_dash, size_pt=13, bold=True)
                r_txt = p_sub.add_run(line)
                set_font(r_txt, size_pt=13, bold=False)
            else:
                r_dash = p.add_run("– ")
                set_font(r_dash, size_pt=13, bold=True)
                r_txt = p.add_run(line)
                set_font(r_txt, size_pt=13, bold=False)

def add_heading_I(doc, text):
    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3, line_spacing=1.15, first_indent=0)
    r = p.add_run(text)
    set_font(r, size_pt=13, bold=True)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=4, space_after=2, line_spacing=1.15, first_indent=0)
    r = p.add_run(text)
    set_font(r, size_pt=13, bold=True)

def add_body_p(doc, text, bold_prefix="", italic_prefix=""):
    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=3, line_spacing=1.15, first_indent=1.25)
    if bold_prefix:
        r_bp = p.add_run(bold_prefix)
        set_font(r_bp, size_pt=13, bold=True)
    if italic_prefix:
        r_ip = p.add_run(italic_prefix)
        set_font(r_ip, size_pt=13, italic=True)
    r_txt = p.add_run(text)
    set_font(r_txt, size_pt=13, bold=False)

def add_bullet_p(doc, text, bold_prefix="", italic_prefix=""):
    p = doc.add_paragraph()
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2, line_spacing=1.15, first_indent=0, left_indent=0.8)
    r_bullet = p.add_run("• ")
    set_font(r_bullet, size_pt=12, bold=True)
    if bold_prefix:
        r_bp = p.add_run(bold_prefix)
        set_font(r_bp, size_pt=13, bold=True)
    if italic_prefix:
        r_ip = p.add_run(italic_prefix)
        set_font(r_ip, size_pt=13, italic=True)
    r_txt = p.add_run(text)
    set_font(r_txt, size_pt=13, bold=False)

# ==================== NỘI DUNG VĂN BẢN 1 ====================
def build_van_ban_1(doc):
    add_header_table(
        doc,
        org_line1="ỦY BAN NHÂN DÂN PHƯỜNG THỦ DẦU MỘT",
        org_line2="PHÒNG KINH TẾ, HẠ TẦNG VÀ ĐÔ THỊ",
        doc_num="Số:         /BC-PKTHTĐT",
        date_str="Thủ Dầu Một, ngày 19 tháng 10 năm 2026"
    )
    
    add_doc_title(
        doc,
        title="BÁO CÁO THẨM ĐỊNH HỒ SƠ VÀ ĐỀ XUẤT HƯỚNG GIẢI QUYẾT",
        subtitle="(V/v: Giải quyết Đơn phản ánh số đến 2346 ngày 15/10/2026 của bà Nguyễn Thị Mộng Đẹp)"
    )
    
    add_kinh_gui(doc, ["Lãnh đạo Phòng Kinh tế, Hạ tầng và Đô thị phường Thủ Dầu Một."])
    
    add_body_p(
        doc,
        "Thực hiện sự phân công của Lãnh đạo Phòng về việc kiểm tra, xác minh và tham mưu xử lý Đơn phản ánh đề ngày 15/10/2026 của bà Nguyễn Thị Mộng Đẹp (Văn phòng HĐND – UBND phường vào Sổ theo dõi công văn ĐẾN số: 2346 ngày 15/10/2026), chuyên viên phụ trách trật tự xây dựng – đô thị báo cáo kết quả thẩm tra hồ sơ địa chính, phân tích kỹ thuật hiện trường và đề xuất hướng xử lý như sau:"
    )
    
    add_heading_I(doc, "I. THÔNG TIN HÀNH CHÍNH CỦA HỒ SƠ:")
    
    add_heading_1(doc, "1. Người đứng đơn phản ánh:")
    add_bullet_p(doc, "NGUYỄN THỊ MỘNG ĐẸP, sinh ngày: 18/03/1962.", bold_prefix="Bà: ")
    add_bullet_p(doc, "074162000264 do Cục Cảnh sát Quản lý hành chính về trật tự xã hội cấp ngày 25/03/2021.", bold_prefix="Số CCCD: ")
    add_bullet_p(doc, "K396/68 đường Nguyễn Tri Phương, phường Chánh Nghĩa, thành phố Thủ Dầu Một, tỉnh Bình Dương (nay là phường Thủ Dầu Một, Thành phố Hồ Chí Minh).", bold_prefix="Địa chỉ thường trú: ")
    add_bullet_p(doc, "0977.975.248.", bold_prefix="Số điện thoại liên lạc: ")
    add_bullet_p(doc, "Là người sử dụng thửa đất số 225, tờ bản đồ số 20 tọa lạc tại phường Thủ Dầu Một, Thành phố Hồ Chí Minh.")
    
    add_heading_1(doc, "2. Người bị phản ánh:")
    add_bullet_p(doc, "ĐOÀN MINH QUANG, sinh năm: 1978.", bold_prefix="Ông: ")
    add_bullet_p(doc, "Số 254/9 đường Nguyễn Tri Phương, phường Chánh Nghĩa, thành phố Thủ Dầu Một, tỉnh Bình Dương (nay là phường Thủ Dầu Một, Thành phố Hồ Chí Minh).", bold_prefix="Địa chỉ cư trú: ")
    add_bullet_p(doc, "0917.747.837.", bold_prefix="Số điện thoại: ")
    add_bullet_p(doc, "Là người sử dụng thửa đất số 220, tờ bản đồ số 20 (thửa đất liền kề).")
    
    add_heading_1(doc, "3. Nội dung phản ánh và yêu cầu của đương sự:")
    add_body_p(
        doc,
        "Bà Nguyễn Thị Mộng Đẹp phản ánh ông Đoàn Minh Quang tiến hành xây dựng công trình trên thửa đất số 220 và tự ý kéo các đường dây điện vào nhà lên trên mái tôn nhà bà Đẹp (thửa đất số 225), cho rằng tiềm ẩn nguy cơ rò rỉ điện trong mùa mưa bão, xâm phạm quyền sở hữu không gian nhà ở và có thái độ thách thức. Đề nghị chính quyền can thiệp buộc ông Quang tháo gỡ, di dời đường dây điện ra khỏi không gian đất của gia đình bà."
    )
    
    add_heading_I(doc, "II. KẾT QUẢ THẨM TRA HỒ SƠ ĐỊA CHÍNH VÀ GIÁM ĐỊNH KỸ THUẬT HIỆN TRƯỜNG:")
    
    add_heading_1(doc, "1. Về cơ sở pháp lý đất đai và tài sản trên đất của thửa đất số 225:")
    add_body_p(
        doc,
        "Căn cứ Giấy chứng nhận quyền sử dụng đất, quyền sở hữu tài sản gắn liền với đất số AA 08063500 (Số vào sổ cấp GCN: CN19343) do Chi nhánh Văn phòng Đăng ký đất đai Số 23 cấp ngày 07/6/2026 cho bà Nguyễn Thị Mộng Đẹp (bản photo do bà Đẹp nộp kèm theo đơn), xác định:"
    )
    add_bullet_p(
        doc,
        "Thửa đất số 225, tờ bản đồ số 20 có diện tích 277,3 m² (gồm 160,0 m² đất ở tại đô thị và 117,3 m² đất trồng cây lâu năm).",
        italic_prefix="Về nguồn gốc và diện tích: "
    )
    add_bullet_p(
        doc,
        'Tại Mục 3 ghi nhận: "3. Thông tin tài sản gắn liền với đất: -/-" và Mục 5 ghi nhận: "5. Ghi chú: -/-". Do đó, về mặt pháp lý đăng ký tài sản, công trình nhà khung sắt mái tôn hiện hữu trên thửa đất số 225 chưa được hoàn công/chứng nhận quyền sở hữu tài sản gắn liền với đất trên Giấy chứng nhận.',
        italic_prefix="Về thông tin công trình xây dựng: "
    )
    add_bullet_p(
        doc,
        'Tại Mục 4 (Sơ đồ thửa đất) thể hiện cạnh ranh giới phía Nam từ điểm ranh số 3 đến điểm ranh số 2 có chiều dài 36,34m tiếp giáp trực tiếp với thửa đất số 220 – Đoàn Minh Quang. Phía Đông Nam (các điểm ranh số 2 – 6 – 5) tiếp giáp trực tiếp với "Lối đi 2.5m" nối thông ra "Đường bê tông".',
        italic_prefix="Về ranh giới tiếp giáp và hạ tầng kỹ thuật: "
    )
    
    add_heading_1(doc, "2. Về giám định kỹ thuật hiện trạng đối với 05 bản ảnh do bà Đẹp cung cấp:")
    add_body_p(
        doc,
        "Qua đối chiếu kỹ thuật đối với 05 trang hình ảnh kèm theo đơn số 2346, chuyên viên phân loại hệ thống dây dẫn thành hai nhóm độc lập:"
    )
    add_bullet_p(
        doc,
        "Là đường dây điện bọc cách điện của hộ ông Đoàn Minh Quang được cố định trên trụ đỡ bằng thanh sắt thẳng đứng gắn trên đỉnh tường rào gạch ranh giới, đầu trụ có lắp sứ cách điện chuyên dụng và co nhựa bảo vệ đầu dây. Hệ thống này chạy dọc theo ranh giới tường rào và không gian lối đi chung 2,5m, không tiếp xúc và không đè lên mặt mái tôn nhà bà Đẹp.",
        italic_prefix="Nhóm 1 – Đường dây tải điện phục vụ sinh hoạt (Trang 5 hồ sơ đơn): "
    )
    add_bullet_p(
        doc,
        "Các sợi dây cáp màu đen đi vắt ngang trên khoảng không và cuộn cáp treo trên thân trụ bê tông sát cổng rào lưới B40 có khóa xích là hệ thống cáp tín hiệu viễn thông của 03 nhà mạng (Viettel, VNPT, Mobifone) kết nối trạm BTS lắp đặt từ năm 2009 theo hợp đồng thuê mặt bằng của đơn vị quản lý trạm (Công ty Toàn Cầu). Đây là cáp tín hiệu viễn thông, không phải dây dẫn điện năng sinh hoạt nên không phát sinh nguy cơ rò rỉ điện giật như nội dung bà Đẹp lo ngại.",
        italic_prefix="Nhóm 2 – Tuyến cáp tín hiệu viễn thông (Trang 3, 4, 6 và 7 hồ sơ đơn): "
    )
    
    add_heading_I(doc, "III. ĐỀ XUẤT PHƯƠNG ÁN GIẢI QUYẾT:")
    add_bullet_p(doc, "Phối hợp cùng Tổ Quản lý trật tự đô thị và Ban điều hành Khu phố Chánh Nghĩa 8 tổ chức kiểm tra thực địa, lập Biên bản kiểm tra, xác minh hiện trường ghi nhận ranh giới sử dụng đất theo Giấy chứng nhận số AA 08063500 cấp ngày 07/6/2026.")
    add_bullet_p(doc, "Tiếp nhận Đơn giải trình của ông Đoàn Minh Quang và văn bản gửi Công ty Toàn Cầu yêu cầu 03 nhà mạng tháo dỡ cáp trạm BTS.")
    add_bullet_p(doc, "Tham mưu Lãnh đạo Phòng ký ban hành Công văn trả lời và giải quyết Đơn phản ánh số 2346 gửi cho bà Nguyễn Thị Mộng Đẹp và ông Đoàn Minh Quang, bác bỏ yêu cầu buộc tháo dỡ đường dây điện sinh hoạt hợp pháp của ông Quang, đồng thời hướng dẫn lộ trình tháo dỡ cáp viễn thông của trạm BTS khi kết thúc hợp đồng thuê mặt bằng để đóng hồ sơ theo luật định.")
    
    # Signature table
    p_sp = doc.add_paragraph()
    set_para_format(p_sp, space_before=4, space_after=2, line_spacing=1.0)
    
    table_sig = doc.add_table(rows=1, cols=2)
    table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table_sig)
    
    col_w_left = Cm(8.0)
    col_w_right = Cm(8.5)
    row_sig = table_sig.rows[0]
    c_left = row_sig.cells[0]
    c_right = row_sig.cells[1]
    c_left.width = col_w_left
    c_right.width = col_w_right
    set_cell_margins(c_left, top=0, bottom=20, left=0, right=10)
    set_cell_margins(c_right, top=0, bottom=20, left=10, right=0)
    
    # Cột trái: Ý kiến Lãnh đạo
    p_ld = c_left.paragraphs[0]
    set_para_format(p_ld, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, line_spacing=1.15)
    r_ld = p_ld.add_run("Ý KIẾN PHÊ DUYỆT CỦA LÃNH ĐẠO PHÒNG\n\n\n\n\n\n")
    set_font(r_ld, size_pt=12, bold=True)
    r_ld_dot = p_ld.add_run("...........................................................................")
    set_font(r_ld_dot, size_pt=12, bold=False)
    
    # Cột phải: Chuyên viên thụ lý
    p_cv = c_right.paragraphs[0]
    set_para_format(p_cv, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, line_spacing=1.15)
    r_cv1 = p_cv.add_run("CHUYÊN VIÊN THỤ LÝ HỒ SƠ\n")
    set_font(r_cv1, size_pt=12, bold=True)
    r_cv2 = p_cv.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    set_font(r_cv2, size_pt=11, italic=True)
    r_cv_dot = p_cv.add_run("...........................................................................")
    set_font(r_cv_dot, size_pt=12, bold=False)

# ==================== NỘI DUNG VĂN BẢN 2 ====================
def build_van_ban_2(doc):
    add_header_table(
        doc,
        org_line1="",
        org_line2="",
        doc_num="",
        date_str="Thủ Dầu Một, ngày 19 tháng 10 năm 2026"
    )
    
    add_doc_title(
        doc,
        title="ĐƠN GIẢI TRÌNH VÀ PHẢN BÁC NỘI DUNG ĐƠN PHẢN ÁNH",
        subtitle="(V/v: Giải trình làm rõ và bác bỏ toàn bộ nội dung Đơn phản ánh số đến 2346 ngày 15/10/2026 của bà Nguyễn Thị Mộng Đẹp)"
    )
    
    add_kinh_gui(doc, [
        "ỦY BAN NHÂN DÂN PHƯỜNG THỦ DẦU MỘT, THÀNH PHỐ HỒ CHÍ MINH;",
        "PHÒNG KINH TẾ, HẠ TẦNG VÀ ĐÔ THỊ PHƯỜNG THỦ DẦU MỘT."
    ])
    
    add_heading_I(doc, "I. THÔNG TIN NGƯỜI LÀM ĐƠN GIẢI TRÌNH:")
    add_bullet_p(doc, "ĐOÀN MINH QUANG                                     Giới tính: Nam", bold_prefix="Họ và tên: ")
    add_bullet_p(doc, "16/04/1978", bold_prefix="Ngày, tháng, năm sinh: ")
    add_bullet_p(doc, "074078003477 cấp ngày 10/3/2023 tại Cục Cảnh sát Quản lý hành chính về trật tự xã hội.", bold_prefix="Số CCCD: ")
    add_bullet_p(doc, "Số 254/9 đường Nguyễn Tri Phương, khu phố Chánh Nghĩa 8, phường Thủ Dầu Một, Thành phố Hồ Chí Minh.", bold_prefix="Nơi đăng ký thường trú: ")
    add_bullet_p(doc, "0917.747.837.", bold_prefix="Số điện thoại liên hệ: ")
    add_bullet_p(doc, "Là người sử dụng đất hợp pháp đối với thửa đất số 220, tờ bản đồ số 20 tọa lạc tại phường Thủ Dầu Một, Thành phố Hồ Chí Minh.", bold_prefix="Quyền lợi và nghĩa vụ liên quan: ")
    
    add_heading_I(doc, "II. NỘI DUNG GIẢI TRÌNH VÀ PHẢN BÁC:")
    add_body_p(
        doc,
        "Ngày 15/10/2026, bà Nguyễn Thị Mộng Đẹp (chủ sử dụng thửa đất số 225, tờ bản đồ số 20 liền kề) đã gửi Đơn phản ánh số đến 2346 đến UBND phường Thủ Dầu Một. Trong đơn, bà Đẹp phản ánh sai sự thật rằng tôi tự ý kéo các đường dây điện vào nhà lên mái tôn nhà bà, gây nguy cơ rò rỉ điện đe dọa tính mạng và có thái độ thách thức, hăm dọa bà. Để bảo vệ quyền và lợi ích hợp pháp của mình, tôi xin giải trình cụ thể các điểm sau:"
    )
    
    add_heading_1(doc, "1. Về đường dây tải điện sinh hoạt của gia đình tôi:")
    add_bullet_p(
        doc,
        "Đường dây điện phục vụ sinh hoạt của tôi hoàn toàn nằm trong phạm vi không gian Lối đi 2.5m nối thông ra Đường bê tông (lối đi chung hợp pháp được thể hiện rõ trên Sơ đồ thửa đất tại Mục 4 Giấy chứng nhận quyền sử dụng đất số AA 08063500 cấp ngày 07/6/2026 của chính bà Nguyễn Thị Mộng Đẹp) và đi dọc theo ranh giới tường rào thuộc thửa đất số 220 của tôi."
    )
    add_bullet_p(
        doc,
        "Tuyến dây điện này được gắn cố định trên trụ đỡ sắt có sứ cách điện chuyên dụng và co nhựa bảo vệ đầu dây (như thể hiện tại bản ảnh trang 5 của bà Đẹp), đảm bảo an toàn tuyệt đối theo Điều 175 và Điều 255 Bộ luật Dân sự năm 2015, hoàn toàn không đè lên mái tôn nhà bà Đẹp."
    )
    
    add_heading_1(doc, "2. Về các sợi dây cáp viễn thông đi phía trên mái tôn trong ảnh chụp của bà Đẹp:")
    add_bullet_p(
        doc,
        "Các sợi dây cáp màu đen đi trên không gian cao và cuộn cáp buộc cố định trên thân trụ bê tông cạnh cổng rào lưới B40 mà bà Đẹp chụp ảnh nộp kèm đơn không phải là dây điện sinh hoạt của tôi mới kéo, mà là hệ thống cáp tín hiệu viễn thông của 03 nhà mạng (Viettel, VNPT, Mobifone) kết nối trạm BTS do Công ty Toàn Cầu thuê mặt bằng từ năm 2009."
    )
    add_bullet_p(
        doc,
        "Tuyến cáp này đã tồn tại từ năm 2009 (trước khi bà Đẹp dựng khung sắt mái tôn sát ranh đất và trước thời điểm bà Đẹp được cấp Giấy chứng nhận số AA 08063500 ngày 07/6/2026). Cáp viễn thông không truyền tải điện năng nên không có nguy cơ rò rỉ điện giật như nội dung đơn đã vu khống."
    )
    add_bullet_p(
        doc,
        "Hiện nay hợp đồng thuê trạm BTS sắp kết thúc (việc chồng bà Đoàn Thị Quý tự ý thỏa thuận miệng gia hạn từ tháng 7 đến tháng 10/2026 và thu giữ 30.000.000 đồng không có sự đồng ý của tôi là hành vi trái pháp luật). Tôi đã chính thức gửi văn bản yêu cầu Công ty Toàn Cầu cùng 03 nhà mạng cử nhân viên kỹ thuật xuống thu hồi, tháo dỡ toàn bộ hệ thống cáp này."
    )
    
    add_heading_1(doc, "3. Về lời cáo buộc \"thách thức, hăm dọa\" và hiện trạng công trình trên thửa đất số 225:")
    add_bullet_p(
        doc,
        "Tôi khẳng định không có hành vi thách thức hay hăm dọa bà Đẹp. Mọi việc tôi đều giải quyết bằng văn bản hành chính với cơ quan chức năng."
    )
    add_bullet_p(
        doc,
        'Ngược lại, tại Mục 3 Giấy chứng nhận số AA 08063500 cấp ngày 07/6/2026 của bà Đẹp ghi rõ: "Thông tin tài sản gắn liền với đất: -/-", nhưng thực tế hiện trường bà Đẹp đã dựng khung sắt mái tôn áp sát vào tường rào ranh giới giáp thửa 220 của tôi suốt chiều dài 36,34m.'
    )
    
    add_heading_I(doc, "III. KIẾN NGHỊ:")
    add_body_p(doc, "Kính đề nghị UBND phường Thủ Dầu Một và Phòng Kinh tế, Hạ tầng và Đô thị:")
    add_bullet_p(doc, "Bác bỏ nội dung Đơn phản ánh số đến 2346 ngày 15/10/2026 của bà Nguyễn Thị Mộng Đẹp đối với đường dây điện sinh hoạt hợp pháp của tôi.")
    add_bullet_p(doc, "Xác nhận hệ thống cáp trên cao là cáp viễn thông của 03 nhà mạng từ năm 2009 và hỗ trợ đôn đốc Công ty Toàn Cầu cùng các nhà mạng tháo dỡ khi thanh lý hợp đồng.")
    add_bullet_p(doc, "Nhắc nhở hộ bà Nguyễn Thị Mộng Đẹp quản lý công trình mái tôn trên thửa đất số 225 đúng ranh giới, không để nước mưa chảy tràn sang ranh đất thửa 220 của gia đình tôi.")
    
    # Signature table
    p_sp = doc.add_paragraph()
    set_para_format(p_sp, space_before=4, space_after=2, line_spacing=1.0)
    
    table_sig = doc.add_table(rows=1, cols=2)
    table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table_sig)
    
    col_w_left = Cm(8.0)
    col_w_right = Cm(8.5)
    row_sig = table_sig.rows[0]
    c_left = row_sig.cells[0]
    c_right = row_sig.cells[1]
    c_left.width = col_w_left
    c_right.width = col_w_right
    set_cell_margins(c_left, top=0, bottom=20, left=0, right=10)
    set_cell_margins(c_right, top=0, bottom=20, left=10, right=0)
    
    # Cột trái: Tài liệu kèm theo
    p_tl = c_left.paragraphs[0]
    set_para_format(p_tl, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2, line_spacing=1.15)
    r_tl_t = p_tl.add_run("Tài liệu kèm theo:\n")
    set_font(r_tl_t, size_pt=12, bold=True)
    r_tl_1 = p_tl.add_run("– Bản sao Hợp đồng thuê trạm BTS (2009);\n– Thông báo yêu cầu tháo dỡ cáp viễn thông;\n– Bản sao CCCD của ông Đoàn Minh Quang.")
    set_font(r_tl_1, size_pt=11, bold=False)
    
    # Cột phải: Người làm đơn giải trình
    p_ng = c_right.paragraphs[0]
    set_para_format(p_ng, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, line_spacing=1.15)
    r_ng1 = p_ng.add_run("NGƯỜI LÀM ĐƠN GIẢI TRÌNH\n")
    set_font(r_ng1, size_pt=12, bold=True)
    r_ng2 = p_ng.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    set_font(r_ng2, size_pt=11, italic=True)
    r_ng3 = p_ng.add_run("Đoàn Minh Quang")
    set_font(r_ng3, size_pt=13, bold=True)

# ==================== NỘI DUNG VĂN BẢN 3 ====================
def build_van_ban_3(doc):
    add_header_table(
        doc,
        org_line1="ỦY BAN NHÂN DÂN PHƯỜNG THỦ DẦU MỘT",
        org_line2="PHÒNG KINH TẾ, HẠ TẦNG VÀ ĐÔ THỊ",
        doc_num="",
        date_str=""
    )
    
    add_doc_title(
        doc,
        title="BIÊN BẢN KIỂM TRA, XÁC MINH HIỆN TRƯỜNG",
        subtitle="(V/v: Kiểm tra thực địa giải quyết Đơn phản ánh số đến 2346 ngày 15/10/2026 của bà Nguyễn Thị Mộng Đẹp)"
    )
    
    add_body_p(
        doc,
        "Hôm nay, vào lúc      giờ      phút, ngày 19 tháng 10 năm 2026, tại thực địa khu vực ranh giới giữa thửa đất số 225 và thửa đất số 220, tờ bản đồ số 20, đường Nguyễn Tri Phương, khu phố Chánh Nghĩa 8, phường Thủ Dầu Một, Thành phố Hồ Chí Minh."
    )
    
    add_heading_I(doc, "I. THÀNH PHẦN THAM GIA ĐOÀN KIỂM TRA:")
    
    add_heading_1(doc, "1. Đại diện Phòng Kinh tế, Hạ tầng và Đô thị phường Thủ Dầu Một:")
    add_bullet_p(doc, "......................................................................... Chức vụ: Lãnh đạo Phòng.", bold_prefix="Ông/Bà: ")
    add_bullet_p(doc, "......................................................................... Chức vụ: Chuyên viên thụ lý hồ sơ.", bold_prefix="Ông/Bà: ")
    add_bullet_p(doc, "......................................................................... Chức vụ: Cán bộ Đội Trật tự đô thị.", bold_prefix="Ông/Bà: ")
    
    add_heading_1(doc, "2. Đại diện Ban điều hành Khu phố Chánh Nghĩa 8:")
    add_bullet_p(doc, "......................................................................... Chức vụ: Trưởng khu phố.", bold_prefix="Ông/Bà: ")
    
    add_heading_1(doc, "3. Các bên đương sự có liên quan:")
    add_bullet_p(doc, "Bà NGUYỄN THỊ MỘNG ĐẸP, sinh năm: 1962; CCCD số: 074162000264 (Chủ sử dụng thửa đất số 225, tờ bản đồ số 20).", bold_prefix="Bên phản ánh: ")
    add_bullet_p(doc, "Ông ĐOÀN MINH QUANG, sinh năm: 1978; CCCD số: 074078003477 (Chủ sử dụng thửa đất số 220, tờ bản đồ số 20).", bold_prefix="Bên bị phản ánh: ")
    
    add_heading_I(doc, "II. KẾT QUẢ KIỂM TRA VÀ ĐỐI CHIẾU THỰC TẾ:")
    add_body_p(
        doc,
        "Đoàn kiểm tra tiến hành đối chiếu Sơ đồ thửa đất tại Mục 4 Giấy chứng nhận quyền sử dụng đất số AA 08063500 cấp ngày 07/6/2026 của thửa 225, hồ sơ thửa đất số 220 và đo vẽ thực địa, lập biên bản ghi nhận các nội dung sau:"
    )
    
    add_heading_1(doc, "1. Về đường dây dẫn điện sinh hoạt của hộ ông Đoàn Minh Quang:")
    add_bullet_p(
        doc,
        'Tại đoạn ranh giới tiếp giáp Lối đi 2.5m nối thông ra Đường bê tông (điểm ranh số 2 – 6 – 5 trên Sơ đồ thửa đất số 225) và dọc theo đỉnh tường rào gạch ranh giới (chiều dài 36,34m): Đường dây điện sinh hoạt của hộ ông Quang là loại cáp bọc cách điện, được gắn trên thanh trụ sắt thẳng đứng có lắp sứ cách điện chuyên dụng và ống co nhựa bảo vệ.'
    )
    add_bullet_p(
        doc,
        'Đo đạc theo phương thẳng đứng từ ranh giới đất: Tuyến dây điện nằm trọn vẹn trong không gian của Lối đi 2.5m và ranh giới tường rào thuộc thửa đất số 220 của ông Quang, không tiếp xúc mặt mái tôn và không xâm phạm không gian thửa đất số 225 của bà Đẹp.'
    )
    
    add_heading_1(doc, "2. Về các sợi dây cáp đi ngang khoảng không mái tôn và cuộn cáp tại trụ bê tông:")
    add_bullet_p(
        doc,
        'Các sợi cáp màu đen đi trên không gian cao phía trên mái tôn và cuộn cáp buộc cố định trên trụ bê tông sát cổng rào lưới B40 được xác định là cáp tín hiệu viễn thông của 03 nhà mạng (Viettel, VNPT, Mobifone) đấu nối trạm BTS đặt trên đất ông Quang từ năm 2009.'
    )
    add_bullet_p(
        doc,
        'Đây là cáp viễn thông, không phải dây dẫn điện năng sinh hoạt mới kéo nên không phát sinh nguy cơ rò rỉ điện. Ông Quang đã có văn bản yêu cầu đơn vị quản lý trạm (Công ty Toàn Cầu) và 03 nhà mạng tháo dỡ thu hồi cáp khi kết thúc hợp đồng thuê mặt bằng.'
    )
    
    add_heading_1(doc, "3. Về hiện trạng công trình trên thửa đất số 225 của bà Nguyễn Thị Mộng Đẹp:")
    add_bullet_p(doc, "Trên thửa 225 hiện hữu công trình nhà khung sắt lợp mái tôn xây dựng áp sát ranh giới tường rào thửa 220.")
    add_bullet_p(doc, 'Đối chiếu Giấy chứng nhận quyền sử dụng đất số AA 08063500 cấp ngày 07/6/2026, tại Mục 3 ghi nhận: "Thông tin tài sản gắn liền với đất: -/-".')
    
    add_heading_I(doc, "III. KẾT LUẬN VÀ KIẾN NGHỊ CỦA ĐOÀN KIỂM TRA:")
    add_bullet_p(doc, "Đường dây tải điện phục vụ sinh hoạt của hộ ông Đoàn Minh Quang nằm trong phạm vi Lối đi chung 2,5m và ranh giới tường rào thửa đất số 220, có sứ cách điện đúng tiêu chuẩn an toàn, không nằm đè lên mái tôn thửa 225.")
    add_bullet_p(doc, "Đối với cáp viễn thông của 03 nhà mạng (Viettel, VNPT, Mobifone) thuộc trạm BTS: Yêu cầu ông Đoàn Minh Quang đôn đốc các đơn vị viễn thông khẩn trương thu hồi, tháo dỡ toàn bộ tuyến cáp khi kết thúc hợp đồng thuê mặt bằng. Đề nghị bà Nguyễn Thị Mộng Đẹp tạo điều kiện thuận lợi để nhân viên kỹ thuật thực hiện thu hồi cáp.")
    add_bullet_p(doc, "Nhắc nhở hai hộ gia đình sử dụng đất đúng ranh giới theo Giấy chứng nhận đã được cơ quan nhà nước có thẩm quyền cấp, giữ gìn trật tự trị an tại địa phương.")
    
    add_body_p(
        doc,
        "Biên bản kết thúc vào lúc      giờ      phút cùng ngày, được lập thành 03 bản có giá trị pháp lý như nhau, đã thông qua cho các bên tham gia cùng nghe và ký tên xác nhận./."
    )
    
    # Signature table 2x2
    p_sp = doc.add_paragraph()
    set_para_format(p_sp, space_before=4, space_after=2, line_spacing=1.0)
    
    table_sig = doc.add_table(rows=2, cols=2)
    table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table_sig)
    
    col_w = Cm(8.25)
    for r in table_sig.rows:
        for c in r.cells:
            c.width = col_w
            set_cell_margins(c, top=0, bottom=15, left=5, right=5)
            
    # Row 0: Hai hộ đương sự
    c00 = table_sig.rows[0].cells[0]
    p00 = c00.paragraphs[0]
    set_para_format(p00, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=1, line_spacing=1.15)
    r00_1 = p00.add_run("ĐẠI DIỆN HỘ BÀ NGUYỄN THỊ MỘNG ĐẸP\n")
    set_font(r00_1, size_pt=12, bold=True)
    r00_2 = p00.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
    set_font(r00_2, size_pt=11, italic=True)
    r00_3 = p00.add_run("Nguyễn Thị Mộng Đẹp\n")
    set_font(r00_3, size_pt=13, bold=True)
    
    c01 = table_sig.rows[0].cells[1]
    p01 = c01.paragraphs[0]
    set_para_format(p01, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=1, line_spacing=1.15)
    r01_1 = p01.add_run("ĐẠI DIỆN HỘ ÔNG ĐOÀN MINH QUANG\n")
    set_font(r01_1, size_pt=12, bold=True)
    r01_2 = p01.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
    set_font(r01_2, size_pt=11, italic=True)
    r01_3 = p01.add_run("Đoàn Minh Quang\n")
    set_font(r01_3, size_pt=13, bold=True)
    
    # Row 1: Khu phố và Phòng
    c10 = table_sig.rows[1].cells[0]
    p10 = c10.paragraphs[0]
    set_para_format(p10, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=1, line_spacing=1.15)
    r10_1 = p10.add_run("ĐẠI DIỆN BAN ĐIỀU HÀNH KHU PHỐ\n")
    set_font(r10_1, size_pt=12, bold=True)
    r10_2 = p10.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
    set_font(r10_2, size_pt=11, italic=True)
    r10_3 = p10.add_run(".........................................................................")
    set_font(r10_3, size_pt=12, bold=False)
    
    c11 = table_sig.rows[1].cells[1]
    p11 = c11.paragraphs[0]
    set_para_format(p11, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=1, line_spacing=1.15)
    r11_1 = p11.add_run("ĐẠI DIỆN PHÒNG KT, HT VÀ ĐÔ THỊ\n")
    set_font(r11_1, size_pt=12, bold=True)
    r11_2 = p11.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
    set_font(r11_2, size_pt=11, italic=True)
    r11_3 = p11.add_run(".........................................................................")
    set_font(r11_3, size_pt=12, bold=False)

# ==================== NỘI DUNG VĂN BẢN 4 ====================
def build_van_ban_4(doc):
    add_header_table(
        doc,
        org_line1="ỦY BAN NHÂN DÂN PHƯỜNG THỦ DẦU MỘT",
        org_line2="PHÒNG KINH TẾ, HẠ TẦNG VÀ ĐÔ THỊ",
        doc_num="Số:         /PKTHTĐT",
        date_str="Thủ Dầu Một, ngày      tháng 10 năm 2026",
        is_cong_van=True,
        trich_yeu_cv="V/v trả lời và giải quyết Đơn phản ánh của bà Nguyễn Thị Mộng Đẹp (Số đến 2346)"
    )
    
    add_kinh_gui(doc, [
        "Bà Nguyễn Thị Mộng Đẹp (Địa chỉ: K396/68 đường Nguyễn Tri Phương, phường Thủ Dầu Một, Thành phố Hồ Chí Minh);",
        "Ông Đoàn Minh Quang (Địa chỉ: 254/9 đường Nguyễn Tri Phương, khu phố Chánh Nghĩa 8, phường Thủ Dầu Một, Thành phố Hồ Chí Minh)."
    ])
    
    add_body_p(
        doc,
        "Ủy ban nhân dân phường Thủ Dầu Một có tiếp nhận Đơn phản ánh đề ngày 15/10/2026 của bà Nguyễn Thị Mộng Đẹp (Văn phòng HĐND – UBND phường vào Sổ công văn ĐẾN số: 2346 ngày 15/10/2026), nội dung phản ánh ông Đoàn Minh Quang (người sử dụng thửa đất số 220, tờ bản đồ số 20) xây dựng công trình và tự ý kéo các đường dây điện vào nhà lên mái tôn nhà bà trên thửa đất số 225, tờ bản đồ số 20, gây nguy cơ rò rỉ điện mùa mưa bão và đề nghị UBND phường can thiệp buộc tháo gỡ, di dời đường dây điện ra khỏi không gian thửa đất số 225."
    )
    
    add_body_p(
        doc,
        "Thực hiện ý kiến chỉ đạo của Lãnh đạo UBND phường Thủ Dầu Một, sau khi thẩm tra hồ sơ địa chính (Giấy chứng nhận quyền sử dụng đất, quyền sở hữu tài sản gắn liền với đất số AA 08063500 do Chi nhánh Văn phòng Đăng ký đất đai Số 23 cấp ngày 07/6/2026 và các tài liệu liên quan) kết hợp kiểm tra thực tế tại hiện trường, Phòng Kinh tế, Hạ tầng và Đô thị thông báo kết quả giải quyết như sau:"
    )
    
    add_heading_1(doc, "1. Đối với đường dây tải điện phục vụ sinh hoạt của hộ ông Đoàn Minh Quang:")
    add_bullet_p(
        doc,
        'Qua kiểm tra thực địa và đối chiếu Sơ đồ thửa đất tại Mục 4 Giấy chứng nhận quyền sử dụng đất số AA 08063500 cấp ngày 07/6/2026, đường dây tải điện sinh hoạt của hộ ông Đoàn Minh Quang là loại cáp bọc cách điện an toàn, được lắp đặt trên trụ đỡ bằng thanh sắt có gắn sứ cách điện chuyên dụng và ống bảo vệ, đi dọc theo phạm vi không gian của Lối đi 2.5m nối thông ra Đường bê tông và ranh giới tường rào thuộc thửa đất số 220 của ông Quang.'
    )
    add_bullet_p(
        doc,
        'Vị trí lắp đặt đường dây điện sinh hoạt nêu trên phù hợp với quy định tại Điều 175 và Điều 255 Bộ luật Dân sự năm 2015, đảm bảo an toàn kỹ thuật, không tiếp xúc mặt mái tôn và không xâm phạm không gian thuộc thửa đất số 225 của bà Nguyễn Thị Mộng Đẹp. Do đó, yêu cầu buộc ông Đoàn Minh Quang tháo gỡ đường dây điện sinh hoạt đang đi trên phạm vi Lối đi 2,5m và ranh giới thửa đất số 220 là không có cơ sở giải quyết.'
    )
    
    add_heading_1(doc, "2. Đối với hệ thống dây cáp đi ngang phía trên khoảng không mái tôn thửa đất số 225:")
    add_bullet_p(
        doc,
        'Các sợi dây cáp màu đen đi trên không gian cao phía trên mái tôn và cuộn cáp buộc trên trụ bê tông sát hàng rào lưới B40 (như hình ảnh đính kèm Đơn phản ánh số 2346) là hệ thống cáp tín hiệu viễn thông của 03 nhà mạng (Viettel, VNPT, Mobifone) đấu nối vào trạm thu phát sóng viễn thông (BTS) đặt trên đất của ông Đoàn Minh Quang từ năm 2009. Đây là cáp tín hiệu viễn thông đã tồn tại từ trước, không phải đường dây tải điện sinh hoạt mới kéo nên không phát sinh nguy cơ rò rỉ điện gây nguy hiểm tính mạng trong mùa mưa bão.'
    )
    add_bullet_p(
        doc,
        'Hiện nay, chủ sử dụng đất thửa số 220 là ông Đoàn Minh Quang đang làm thủ tục chấm dứt hợp đồng cho thuê mặt bằng trạm BTS và đã có văn bản yêu cầu đơn vị quản lý trạm (Công ty Toàn Cầu) cùng 03 đơn vị viễn thông (Viettel, VNPT, Mobifone) cử bộ phận kỹ thuật xuống thu hồi, tháo dỡ toàn bộ hệ thống cáp viễn thông nêu trên.'
    )
    
    add_heading_1(doc, "3. Đề nghị đối với các bên đương sự:")
    add_bullet_p(
        doc,
        'Đề nghị ông Đoàn Minh Quang tiếp tục đôn đốc đơn vị quản lý trạm BTS (Công ty Toàn Cầu) và 03 nhà mạng viễn thông (Viettel, VNPT, Mobifone) khẩn trương thu dọn, tháo dỡ toàn bộ các tuyến cáp viễn thông đi qua khoảng không phía trên mái tôn thửa đất số 225 khi kết thúc hợp đồng thuê mặt bằng.'
    )
    add_bullet_p(
        doc,
        'Đề nghị bà Nguyễn Thị Mộng Đẹp phối hợp, tạo điều kiện thuận lợi khi nhân viên kỹ thuật của các nhà mạng viễn thông đến thực hiện việc thu hồi dây cáp; đồng thời quản lý, sử dụng công trình nhà khung sắt mái tôn trên thửa đất số 225 đúng ranh giới chiều dài 36,34m giáp thửa đất số 220, đảm bảo quy định về thoát nước mưa và an toàn công trình liền kề theo Điều 175, Điều 177 Bộ luật Dân sự năm 2015 và Giấy chứng nhận số AA 08063500 cấp ngày 07/6/2026.'
    )
    add_bullet_p(
        doc,
        'Trong quá trình sinh hoạt tại địa phương, đề nghị hai gia đình giữ gìn trật tự trị an khu phố; trường hợp phát sinh hành vi đe dọa, thách thức gây mất an ninh trật tự, đề nghị công dân thông báo ngay cho Công an phường Thủ Dầu Một để được xử lý theo quy định pháp luật.'
    )
    
    add_body_p(
        doc,
        "Phòng Kinh tế, Hạ tầng và Đô thị phường Thủ Dầu Một thông báo đến bà Nguyễn Thị Mộng Đẹp và ông Đoàn Minh Quang được biết, thực hiện./."
    )
    
    # Signature & Nơi nhận table
    p_sp = doc.add_paragraph()
    set_para_format(p_sp, space_before=4, space_after=2, line_spacing=1.0)
    
    table_sig = doc.add_table(rows=1, cols=2)
    table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table_sig)
    
    col_w_left = Cm(8.0)
    col_w_right = Cm(8.5)
    row_sig = table_sig.rows[0]
    c_left = row_sig.cells[0]
    c_right = row_sig.cells[1]
    c_left.width = col_w_left
    c_right.width = col_w_right
    set_cell_margins(c_left, top=0, bottom=20, left=0, right=10)
    set_cell_margins(c_right, top=0, bottom=20, left=10, right=0)
    
    # Cột trái: Nơi nhận
    p_nn = c_left.paragraphs[0]
    set_para_format(p_nn, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2, line_spacing=1.15)
    r_nn_t = p_nn.add_run("Nơi nhận:\n")
    set_font(r_nn_t, size_pt=12, bold=True, italic=True)
    r_nn_list = p_nn.add_run(
        "– Như trên;\n"
        "– Chủ tịch, các PCT UBND phường (để b/c);\n"
        "– Văn phòng HĐND – UBND phường (để theo dõi);\n"
        "– Công an phường Thủ Dầu Một (để phối hợp);\n"
        "– Ban điều hành Khu phố Chánh Nghĩa 8;\n"
        "– Lưu: VT, PKTHTĐT."
    )
    set_font(r_nn_list, size_pt=11, bold=False)
    
    # Cột phải: Trưởng phòng
    p_tp = c_right.paragraphs[0]
    set_para_format(p_tp, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, line_spacing=1.15)
    r_tp1 = p_tp.add_run("TRƯỞNG PHÒNG\n")
    set_font(r_tp1, size_pt=13, bold=True)
    r_tp2 = p_tp.add_run("(Ký tên, đóng dấu)\n\n\n\n\n")
    set_font(r_tp2, size_pt=11, italic=True)
    r_tp3 = p_tp.add_run(".........................................................................")
    set_font(r_tp3, size_pt=12, bold=False)

# ==================== TRANG BÌA HỒ SƠ CÔNG VỤ ====================
def build_cover_page(doc):
    # Khung viền trang bìa công vụ (Table box bao trọn)
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Đặt viền đôi xanh đậm cho bìa hồ sơ
    tblPr = cover_table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="double" w:sz="18" w:space="0" w:color="1B365D"/>'
        f'<w:left w:val="double" w:sz="18" w:space="0" w:color="1B365D"/>'
        f'<w:bottom w:val="double" w:sz="18" w:space="0" w:color="1B365D"/>'
        f'<w:right w:val="double" w:sz="18" w:space="0" w:color="1B365D"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)
    
    cell = cover_table.rows[0].cells[0]
    cell.width = Cm(16.5)
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)
    
    p0 = cell.paragraphs[0]
    set_para_format(p0, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=2, line_spacing=1.15)
    r0_1 = p0.add_run("ỦY BAN NHÂN DÂN PHƯỜNG THỦ DẦU MỘT\n")
    set_font(r0_1, size_pt=13, bold=False)
    r0_2 = p0.add_run("PHÒNG KINH TẾ, HẠ TẦNG VÀ ĐÔ THỊ\n")
    set_font(r0_2, size_pt=13, bold=True)
    r0_line = p0.add_run("━━━━━━━━━━━━━━━━━━━\n\n\n")
    set_font(r0_line, size_pt=11, bold=True, color_rgb=(27, 54, 93))
    
    # Title
    r_title1 = p0.add_run("HỒ SƠ CÔNG VỤ\n")
    set_font(r_title1, size_pt=24, bold=True, color_rgb=(27, 54, 93))
    r_title2 = p0.add_run("XỬ LÝ ĐƠN PHẢN ÁNH SỐ 2346/UBND\n")
    set_font(r_title2, size_pt=16, bold=True, color_rgb=(200, 30, 30))
    r_title3 = p0.add_run("(V/v: Phản ánh của bà Nguyễn Thị Mộng Đẹp đối với hộ ông Đoàn Minh Quang tại thửa đất số 225 và 220, tờ bản đồ số 20, phường Thủ Dầu Một)\n\n\n")
    set_font(r_title3, size_pt=12, italic=True)
    
    # Bảng thông tin hồ sơ
    p_info = cell.add_paragraph()
    set_para_format(p_info, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=6, line_spacing=1.2, left_indent=0.5)
    r_inf_t = p_info.add_run("THÔNG TIN QUẢN LÝ HỒ SƠ:\n")
    set_font(r_inf_t, size_pt=12, bold=True, color_rgb=(27, 54, 93))
    
    items = [
        ("Mã số hồ sơ:", " 2346/HS-PKTHTĐT"),
        ("Ngày thụ lý:", " 15 tháng 10 năm 2026 (Sổ công văn ĐẾN số 2346)"),
        ("Ngày kết luận giải quyết:", " 19 tháng 10 năm 2026"),
        ("Cơ quan chủ trì giải quyết:", " Phòng Kinh tế, Hạ tầng và Đô thị phường Thủ Dầu Một"),
        ("Cơ quan phối hợp:", " Ban điều hành Khu phố Chánh Nghĩa 8, Đội Trật tự đô thị"),
        ("Chuyên viên thụ lý:", " Bộ phận Quản lý Trật tự Xây dựng – Đô thị"),
        ("Người đứng đơn phản ánh:", " Bà Nguyễn Thị Mộng Đẹp (Thửa đất số 225, TBĐ 20)"),
        ("Người bị phản ánh:", " Ông Đoàn Minh Quang (Thửa đất số 220, TBĐ 20)"),
        ("Tình trạng giải quyết:", " Đã kiểm tra thực địa, lập Biên bản và ban hành Công văn trả lời chính thức."),
    ]
    for lbl, val in items:
        p_item = cell.add_paragraph()
        set_para_format(p_item, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=1, space_after=1, line_spacing=1.15, left_indent=0.8)
        r_lbl = p_item.add_run("• " + lbl)
        set_font(r_lbl, size_pt=11, bold=True)
        r_val = p_item.add_run(val)
        set_font(r_val, size_pt=11, bold=False)
        
    p_ml = cell.add_paragraph()
    set_para_format(p_ml, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4, line_spacing=1.2, left_indent=0.5)
    r_ml_t = p_ml.add_run("DANH MỤC VĂN BẢN TRONG HỒ SƠ:\n")
    set_font(r_ml_t, size_pt=12, bold=True, color_rgb=(27, 54, 93))
    
    docs_list = [
        "VĂN BẢN 1: Báo cáo thẩm định hồ sơ và đề xuất hướng giải quyết của chuyên viên thụ lý (Số: .../BC-PKTHTĐT ngày 19/10/2026)",
        "VĂN BẢN 2: Đơn giải trình và phản bác nội dung đơn phản ánh của ông Đoàn Minh Quang (ngày 19/10/2026)",
        "VĂN BẢN 3: Biên bản kiểm tra, xác minh hiện trường thực địa ngày 19/10/2026 (Đoàn liên ngành 4 bên)",
        "VĂN BẢN 4: Công văn trả lời và giải quyết Đơn phản ánh chính thức của Phòng Kinh tế, Hạ tầng và Đô thị (Số: .../PKTHTĐT)"
    ]
    for d in docs_list:
        p_d = cell.add_paragraph()
        set_para_format(p_d, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=1, space_after=2, line_spacing=1.15, left_indent=0.8)
        r_d = p_d.add_run("1." if d == docs_list[0] else ("2." if d == docs_list[1] else ("3." if d == docs_list[2] else "4.")))
        set_font(r_d, size_pt=11, bold=True)
        r_dt = p_d.add_run(" " + d)
        set_font(r_dt, size_pt=11, bold=False)
        
    p_foot = cell.add_paragraph()
    set_para_format(p_foot, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=10, line_spacing=1.15)
    r_foot = p_foot.add_run("LƯU TRỮ VĂN THƯ NĂM 2026 – THỜI HẠN BẢO QUẢN: LÂU DÀI\nThủ Dầu Một, tháng 10 năm 2026")
    set_font(r_foot, size_pt=11, italic=True, bold=True, color_rgb=(100, 100, 100))

# ==================== MAIN GENERATOR ====================
def main():
    workspace_dir = r"q:\CODE_visual\filebctt"
    
    # 1. TẠO FILE TỔNG HỢP TRỌN BỘ HỒ SƠ
    print("Đang tạo file tổng hợp trọn bộ Hồ sơ công vụ...")
    doc_full = docx.Document()
    set_page_setup(doc_full.sections[0])
    
    # Trang bìa
    build_cover_page(doc_full)
    
    # Văn bản 1
    doc_full.add_page_break()
    build_van_ban_1(doc_full)
    
    # Văn bản 2
    doc_full.add_page_break()
    build_van_ban_2(doc_full)
    
    # Văn bản 3
    doc_full.add_page_break()
    build_van_ban_3(doc_full)
    
    # Văn bản 4
    doc_full.add_page_break()
    build_van_ban_4(doc_full)
    
    full_path = os.path.join(workspace_dir, "HO_SO_CONG_VU_XU_LY_DON_2346_UBND.docx")
    doc_full.save(full_path)
    print(f"-> Đã xuất thành công: {full_path}")
    
    # 2. TẠO CÁC FILE RIÊNG LẺ ĐỂ TIỆN IN ẤN TỪNG VĂN BẢN
    files = [
        ("01_Bao_Cao_Tham_Dinh_2346.docx", build_van_ban_1),
        ("02_Don_Giai_Trinh_Doan_Minh_Quang.docx", build_van_ban_2),
        ("03_Bien_Ban_Kiem_Tra_Hien_Truong.docx", build_van_ban_3),
        ("04_Cong_Van_Tra_Loi_Giai_Quyet_2346.docx", build_van_ban_4)
    ]
    
    for filename, builder in files:
        d = docx.Document()
        set_page_setup(d.sections[0])
        builder(d)
        p = os.path.join(workspace_dir, filename)
        d.save(p)
        print(f"-> Đã xuất file riêng: {p}")

if __name__ == "__main__":
    main()
