import os
import sys
import zipfile
import docx
from docx.oxml.ns import qn

# Ensure utf-8 output
sys.stdout.reconfigure(encoding='utf-8')

docx_path = 'V2_DoamMinhQuan_2224801030137_BCTT (2).docx'
doc = docx.Document(docx_path)

output_dir = os.path.join('docs', 'v2_images')
os.makedirs(output_dir, exist_ok=True)

# 1. Extract all media from zip
with zipfile.ZipFile(docx_path, 'r') as z:
    for filename in z.namelist():
        if filename.startswith('word/media/'):
            basename = os.path.basename(filename)
            target = os.path.join(output_dir, basename)
            with open(target, 'wb') as f:
                f.write(z.read(filename))
            print(f"Extracted: {filename} -> {target} ({os.path.getsize(target)} bytes)")

# 2. Map rId to target in doc.part.rels
rel_map = {}
for rel_id, rel in doc.part.rels.items():
    if "image" in rel.target_ref:
        rel_map[rel_id] = os.path.basename(rel.target_ref)
print("\nRelationship Map (rId -> image):")
for k, v in rel_map.items():
    print(f"  {k} -> {v}")

# 3. Traverse document paragraphs and find images and captions
print("\n=== DOCUMENT STRUCTURE & IMAGES ===")
current_chapter = "MỞ ĐẦU / KHÁC"

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    
    # Check heading/chapter
    if any(ch in txt.upper() for ch in ["CHƯƠNG 1", "CHƯƠNG I", "CHƯƠNG 2", "CHƯƠNG II", "CHƯƠNG 3", "CHƯƠNG III", "CHƯƠNG 4", "CHƯƠNG IV"]):
        current_chapter = txt
        print(f"\n--- [P{i:03d}] {current_chapter} ---")
    elif p.style.name.startswith("Heading") or (len(txt) > 0 and len(txt) < 80 and txt.isupper()):
        # Possible section
        pass

    # Check blip r:embed in paragraph xml
    blips = p._element.xpath('.//a:blip')
    if blips:
        for b in blips:
            embed_id = b.get(qn('r:embed'))
            img_file = rel_map.get(embed_id, 'UNKNOWN')
            
            # Look at neighboring paragraphs for caption
            prev_txt = doc.paragraphs[i-1].text.strip() if i > 0 else ""
            next_txt = doc.paragraphs[i+1].text.strip() if i + 1 < len(doc.paragraphs) else ""
            next2_txt = doc.paragraphs[i+2].text.strip() if i + 2 < len(doc.paragraphs) else ""
            
            print(f"\n>>> IMAGE FOUND at P{i:03d} in [{current_chapter}]:")
            print(f"    rId: {embed_id} -> {img_file}")
            print(f"    Prev text: {prev_txt}")
            print(f"    Curr text: {txt}")
            print(f"    Next text: {next_txt}")
            print(f"    Next+1 text: {next2_txt}")

# Check tables as well
for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            blips = cell._tc.xpath('.//a:blip')
            if blips:
                for b in blips:
                    embed_id = b.get(qn('r:embed'))
                    img_file = rel_map.get(embed_id, 'UNKNOWN')
                    print(f"\n>>> IMAGE IN TABLE {t_idx} [R{r_idx}, C{c_idx}]: {embed_id} -> {img_file}")
