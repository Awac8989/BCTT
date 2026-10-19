import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_page_borders(section, color="002060", val="double", sz="12", space="20"):
    sectPr = section._sectPr
    # Remove existing pgBorders if any
    existing = sectPr.find(qn('w:pgBorders'))
    if existing is not None:
        sectPr.remove(existing)
        
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for side in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), val)
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:space'), str(space))
        b.set(qn('w:color'), color)
        pgBorders.append(b)
    sectPr.append(pgBorders)

doc = docx.Document()

# Section 1: Bìa cứng ngoài
s1 = doc.sections[0]
s1.top_margin = Cm(2.0)
s1.bottom_margin = Cm(2.0)
s1.left_margin = Cm(3.0)
s1.right_margin = Cm(2.0)
set_page_borders(s1, color="002060", val="double", sz="18", space="20")

p1 = doc.add_paragraph("UBND TP HỒ CHÍ MINH\nTRƯỜNG ĐẠI HỌC THỦ DẦU MỘT")
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p1.runs[0].font.name = "Times New Roman"
p1.runs[0].font.size = Pt(13)
p1.runs[0].font.bold = True

doc.add_paragraph("\n\nBÁO CÁO CHUYÊN ĐỀ\nTHỰC TẬP DOANH NGHIỆP\n\n").alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Bìa cứng bên ngoài (Màu xanh dương)").alignment = WD_ALIGN_PARAGRAPH.CENTER

# Section 2: Bìa phụ
s2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
s2.top_margin = Cm(2.0)
s2.bottom_margin = Cm(2.0)
s2.left_margin = Cm(3.0)
s2.right_margin = Cm(2.0)
set_page_borders(s2, color="000000", val="double", sz="18", space="20")

p2 = doc.add_paragraph("UBND TP HỒ CHÍ MINH\nTRƯỜNG ĐẠI HỌC THỦ DẦU MỘT")
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.runs[0].font.name = "Times New Roman"
p2.runs[0].font.size = Pt(13)
p2.runs[0].font.bold = True
doc.add_paragraph("\n\nBÁO CÁO CHUYÊN ĐỀ\nTHỰC TẬP DOANH NGHIỆP\n\n").alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Bìa phụ bên trong (Màu đen)").alignment = WD_ALIGN_PARAGRAPH.CENTER

# Section 3: Nội dung chính không có border
s3 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
s3.top_margin = Cm(2.0)
s3.bottom_margin = Cm(2.0)
s3.left_margin = Cm(3.0)
s3.right_margin = Cm(2.0)
# Ensure s3 has no pgBorders
sectPr3 = s3._sectPr
existing = sectPr3.find(qn('w:pgBorders'))
if existing is not None:
    sectPr3.remove(existing)
else:
    # Set none explicitly
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for side in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'none')
        pgBorders.append(b)
    sectPr3.append(pgBorders)

p3 = doc.add_paragraph("LỜI CẢM ƠN")
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Đây là trang nội dung, không có viền bao quanh.")

doc.save("test_borders.docx")
print("Saved test_borders.docx successfully")
