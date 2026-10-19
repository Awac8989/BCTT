import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('BAO_CAO_THUC_TAP_100_TRANG_DOAN_MINH_QUAN.docx')

print("=== 1. CHECK CODE PARAGRAPHS IN BODY TEXT ===")
code_paras = []
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if r.font.name == 'Consolas':
            code_paras.append((i, p.text.strip()))
            break

print(f"Total body paragraphs with Consolas: {len(code_paras)}")
for idx, txt in code_paras:
    print(f"  P{idx}: {txt[:80]}")

print("\n=== 2. CHECK ALL FIGURES IN DOCUMENT ===")
img_captions = []
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("Hình "):
        img_captions.append((i, t))

print(f"Total figures found with caption 'Hình ...': {len(img_captions)}")
for idx, cap in img_captions:
    print(f"  P{idx:03d}: {cap}")

print("\n=== 3. CHECK SECTIONS 4.1.8, 4.1.9, 4.1.10, PHU LUC 2 ===")
targets = ["4.1.8", "4.1.9", "4.1.10", "PHỤ LỤC 2"]
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if any(tg in t for tg in targets):
        print(f"\n[P{i:03d}] {t}")
        for next_i in range(i+1, min(i+4, len(doc.paragraphs))):
            next_t = doc.paragraphs[next_i].text.strip()
            if next_t:
                print(f"   -> P{next_i:03d}: {next_t[:100]}...")

print("\n=== 4. SUMMARY DOCUMENT METRICS ===")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")
