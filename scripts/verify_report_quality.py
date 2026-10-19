import docx
import sys
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

file_path = r"Q:\CODE_visual\filebctt\BAO_CAO_THUC_TAP_100_TRANG_DOAN_MINH_QUAN.docx"
doc = docx.Document(file_path)

print(f"=== KIỂM TRA CHẤT LƯỢNG TÀI LIỆU WORD: {file_path} ===")
print(f"Tổng số đoạn văn (paragraphs): {len(doc.paragraphs)}")
print(f"Tổng số bảng (tables): {len(doc.tables)}")

# 1. Kiểm tra Bookmarks trong tài liệu
bookmarks = []
for p in doc.paragraphs:
    bms = p._p.xpath('.//w:bookmarkStart')
    for b in bms:
        name = b.get(qn('w:name'))
        if not name.startswith('_'):  # Loại bỏ các bookmark ẩn tự động của Word
            bookmarks.append(name)

fig_bms = [b for b in bookmarks if b.startswith('fig_')]
ref_bms = [b for b in bookmarks if b.startswith('ref_')]
tbl_bms = [b for b in bookmarks if b.startswith('tbl_')]
sec_bms = [b for b in bookmarks if b.startswith('sec_')]

print(f"\n--- 1. BOOKMARKS ({len(bookmarks)} bookmarks) ---")
print(f"  - Bookmarks Hình ảnh (fig_*): {len(fig_bms)} (mục tiêu 30 hình)")
print(f"    Ví dụ: {fig_bms[:5]} ... {fig_bms[-3:]}")
print(f"  - Bookmarks Tài liệu tham khảo (ref_*): {len(ref_bms)} (mục tiêu 20 tài liệu)")
print(f"    Danh sách: {ref_bms}")
print(f"  - Bookmarks Bảng biểu (tbl_*): {len(tbl_bms)}")
print(f"    Ví dụ: {tbl_bms[:5]} ... {tbl_bms[-3:]}")
print(f"  - Bookmarks Đề mục (sec_*): {len(sec_bms)}")
print(f"    Danh sách: {sec_bms}")

# 2. Kiểm tra Hyperlinks
hyperlinks = []
for p in doc.paragraphs:
    hls = p._p.xpath('.//w:hyperlink')
    for h in hls:
        anchor = h.get(qn('w:anchor'))
        text = "".join(h.xpath('.//w:t/text()'))
        hyperlinks.append((text, anchor))

fig_links = [h for h in hyperlinks if h[1] and h[1].startswith('fig_')]
ref_links = [h for h in hyperlinks if h[1] and h[1].startswith('ref_')]
tbl_links = [h for h in hyperlinks if h[1] and h[1].startswith('tbl_')]
sec_links = [h for h in hyperlinks if h[1] and h[1].startswith('sec_')]

print(f"\n--- 2. HYPERLINKS NỘI BỘ ({len(hyperlinks)} hyperlinks) ---")
print(f"  - Hyperlinks từ Danh mục Hình ảnh trỏ đến hình: {len(fig_links)}")
for f_txt, f_anc in fig_links[:5]:
    print(f"    '{f_txt[:45]}' -> {f_anc}")
print(f"  - Hyperlinks từ câu văn bản trỏ đến Tài liệu tham khảo: {len(ref_links)}")
for r_txt, r_anc in ref_links[:10]:
    print(f"    '{r_txt}' -> {r_anc}")
print(f"  - Hyperlinks từ Danh mục Bảng biểu trỏ đến bảng: {len(tbl_links)}")
for t_txt, t_anc in tbl_links[:5]:
    print(f"    '{t_txt[:45]}' -> {t_anc}")
print(f"  - Hyperlinks từ Mục lục chính trỏ đến các chương: {len(sec_links)}")

# 3. Kiểm tra Indentation và Spacing
print("\n--- 3. KIỂM TRA THỤT ĐẦU DÒNG VÀ KHOẢNG CÁCH ---")
body_paras = []
headings = []
captions = []
lists = []

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    pf = p.paragraph_format
    fl = pf.first_line_indent.cm if pf.first_line_indent is not None else None
    li = pf.left_indent.cm if pf.left_indent is not None else None
    
    if p.style.name.startswith('Heading') or any(txt.startswith(s) for s in ["CHƯƠNG", "MỞ ĐẦU", "LỜI CẢM ƠN", "MỤC LỤC", "DANH MỤC"]):
        headings.append((txt[:40], fl, li, p.alignment))
    elif txt.startswith("Hình ") or txt.startswith("Bảng "):
        captions.append((txt[:40], fl, li, p.alignment))
    elif txt.startswith("-") or txt.startswith("+") or txt.startswith("•"):
        lists.append((txt[:40], fl, li, p.alignment))
    else:
        body_paras.append((txt[:40], fl, li, p.alignment))

print(f"  - Số đoạn văn nội dung thông thường: {len(body_paras)}")
print(f"    Mẫu 3 đoạn: {body_paras[:3]}")
print(f"  - Số tiêu đề: {len(headings)}")
print(f"    Mẫu 3 tiêu đề: {headings[:3]}")
print(f"  - Số chú thích (Caption): {len(captions)}")
print(f"    Mẫu 3 chú thích: {captions[:3]}")
print(f"  - Số mục danh sách (List items): {len(lists)}")
print(f"    Mẫu 3 mục danh sách: {lists[:3]}")

# 4. Kiểm tra phần Tài liệu tham khảo
print("\n--- 4. KIỂM TRA MỤC TÀI LIỆU THAM KHẢO ---")
ref_paras = [p for p in doc.paragraphs if p.text.strip().startswith('[') and any(f"[{i}]" in p.text for i in range(1, 21))]
print(f"  - Tổng số mục tham khảo được định dạng ở cuối bài: {len(ref_paras)}")
for p in ref_paras[:5]:
    bm = [b.get(qn('w:name')) for b in p._p.xpath('.//w:bookmarkStart')]
    print(f"    {p.text[:60]} | Bookmarks: {bm}")

print("\nKIỂM TRA HOÀN TẤT!")
