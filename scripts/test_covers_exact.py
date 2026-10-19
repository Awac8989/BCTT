import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_page_borders(section, color="002060", val="double", sz="18", space="20"):
    sectPr = section._sectPr
    existing = sectPr.find(qn('w:pgBorders'))
    if existing is not None:
        sectPr.remove(existing)
        
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for side in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), val)
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:space'), str(space))
        b.set(qn('w:color'), color)
        pgBorders.append(b)
    sectPr.append(pgBorders)

def add_cover_content(doc, is_outer=True):
    # 1. Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(10)
    p_header.paragraph_format.space_after = Pt(4)
    p_header.paragraph_format.line_spacing = 1.15
    
    r1 = p_header.add_run("UBND TP HỒ CHÍ MINH\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0)
    
    r2 = p_header.add_run("TRƯỜNG ĐẠI HỌC THỦ DẦU MỘT")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 0, 0)
    
    # 2. Logo
    logo_path = 'docs/v2_images/image1.png'
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(14)
    p_logo.paragraph_format.space_after = Pt(20)
    if os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, width=Cm(4.8))
        
    # 3. Report Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(24)
    p_title.paragraph_format.line_spacing = 1.2
    
    r_t1 = p_title.add_run("BÁO CÁO CHUYÊN ĐỀ\n")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(17)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(0, 0, 0)
    
    r_t2 = p_title.add_run("THỰC TẬP DOANH NGHIỆP")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(17)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(0, 0, 0)
    
    # 4. Location & Subject
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(10)
    p_sub.paragraph_format.space_after = Pt(36)
    p_sub.paragraph_format.line_spacing = 1.25
    
    r_s1 = p_sub.add_run("Nơi thực tập\n")
    r_s1.font.name = "Times New Roman"
    r_s1.font.size = Pt(13)
    r_s1.font.bold = True
    r_s1.font.color.rgb = RGBColor(0, 0, 0)
    
    r_s2 = p_sub.add_run("Phòng Văn hóa- Xã hội phường Thủ Dầu Một\n")
    r_s2.font.name = "Times New Roman"
    r_s2.font.size = Pt(14)
    r_s2.font.bold = True
    r_s2.font.color.rgb = RGBColor(0, 0, 0)
    
    r_s3 = p_sub.add_run("Phân tích, thiết kế và xây dựng hệ thống phục vụ người có công sở lao động thương binh và xã hội")
    r_s3.font.name = "Times New Roman"
    r_s3.font.size = Pt(14)
    r_s3.font.bold = True
    r_s3.font.color.rgb = RGBColor(0, 0, 0)
    
    # 5. Student & Advisor Info block
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Cm(0.5)
    p_info.paragraph_format.space_before = Pt(16)
    p_info.paragraph_format.space_after = Pt(40)
    p_info.paragraph_format.line_spacing = 1.3
    
    info_lines = [
        "NGƯỜI HƯỚNG DẪN TẠI ĐƠN VỊ: HỨA TRỌNG DUY",
        "GIẢNG VIÊN HƯỚNG DẪN: THS. NGUYỄN THANH TRUNG",
        "SINH VIÊN THỰC HIỆN: ĐOÀN MINH QUÂN       MÃ SỐ SV: 2224801030137",
        "CHUYÊN NGÀNH: KỸ THUẬT PHẦN MỀM LỚP: D22KTPM03",
        "NIÊN KHÓA: 2022 - 2027"
    ]
    for idx, line in enumerate(info_lines):
        r = p_info.add_run(line + ("\n" if idx < len(info_lines) - 1 else ""))
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        
    # 6. Footer (City & Year)
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.paragraph_format.space_before = Pt(20)
    p_footer.paragraph_format.space_after = Pt(10)
    
    r_f = p_footer.add_run("TP. Hồ Chí Minh, năm 2026")
    r_f.font.name = "Times New Roman"
    r_f.font.size = Pt(13)
    r_f.font.bold = True
    r_f.font.italic = True
    r_f.font.color.rgb = RGBColor(0, 0, 0)

doc = docx.Document()

# Section 1: Bìa cứng ngoài
s1 = doc.sections[0]
s1.top_margin = Cm(2.0)
s1.bottom_margin = Cm(2.0)
s1.left_margin = Cm(3.0)
s1.right_margin = Cm(2.0)
set_page_borders(s1, color="002060", val="double", sz="18", space="20")
add_cover_content(doc, is_outer=True)

# Section 2: Bìa phụ trong
s2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
s2.top_margin = Cm(2.0)
s2.bottom_margin = Cm(2.0)
s2.left_margin = Cm(3.0)
s2.right_margin = Cm(2.0)
set_page_borders(s2, color="000000", val="double", sz="18", space="20")
add_cover_content(doc, is_outer=False)

doc.save("test_covers_exact.docx")
print("Saved test_covers_exact.docx")
