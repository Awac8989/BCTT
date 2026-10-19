import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.oxml.ns import qn

file_path = "BAO_CAO_THUC_TAP_DOAN_MINH_QUAN_BAN_IN_CHUAN.docx"
doc = docx.Document(file_path)

print(f"=== KIỂM TRA CHẤT LƯỢNG BẢN IN CHUẨN: {file_path} ===")
print("Tổng số đoạn văn (paragraphs):", len(doc.paragraphs))
print("Tổng số bảng (tables):", len(doc.tables))
print("Tổng số sections:", len(doc.sections))

# 1. Kiểm tra Sections & Borders & Margins
print("\n--- 1. SECTIONS, BORDERS & MARGINS ---")
for i, s in enumerate(doc.sections):
    sectPr = s._sectPr
    pgBorders = sectPr.find(qn('w:pgBorders'))
    border_info = "None"
    if pgBorders is not None:
        borders = [f"{c.tag.split('}')[-1]}: val={c.get(qn('w:val'))}, color={c.get(qn('w:color'))}" for c in pgBorders]
        border_info = "; ".join(borders)
    print(f"Section {i}: Top={s.top_margin.cm:.1f}cm, Bottom={s.bottom_margin.cm:.1f}cm, Left={s.left_margin.cm:.1f}cm, Right={s.right_margin.cm:.1f}cm | Borders: {border_info}")

# 2. Kiểm tra chấm tròn đầu dòng (Bullets)
print("\n--- 2. CHẤM TRÒN ĐẦU DÒNG (BULLETS) ---")
bullet_chars = ['•', '●', '○', '■', '◆', '▪', '\u2022', '\u25cf', '\u25cb', '\u25aa']
bullet_found = 0
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    for bc in bullet_chars:
        if t.startswith(bc):
            bullet_found += 1
            print(f"  Cảnh báo P[{i}]: {t[:60]}")
            break

print(f"Tổng số đoạn văn bắt đầu bằng chấm tròn (yêu cầu = 0): {bullet_found}")

# 3. Kiểm tra màu sắc (Chữ số và chữ thường có màu đen không)
print("\n--- 3. MÀU SẮC CHỮ SỐ VÀ CHỮ VIẾT (YÊU CẦU 100% MÀU ĐEN) ---")
non_black_runs = 0
examples_non_black = []

for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if r.font.color and r.font.color.rgb:
            rgb_str = str(r.font.color.rgb)
            if rgb_str != "000000":
                non_black_runs += 1
                if len(examples_non_black) < 5:
                    examples_non_black.append((i, rgb_str, r.text[:40]))
    # check hyperlinks
    for hl in p._p.xpath('.//w:hyperlink'):
        for r in hl.xpath('.//w:r'):
            color_nodes = r.xpath('.//w:rPr/w:color')
            for cn in color_nodes:
                c_val = cn.get(qn('w:val'))
                if c_val and c_val != "000000":
                    non_black_runs += 1
                    if len(examples_non_black) < 5:
                        examples_non_black.append((i, f"hl_color={c_val}", "".join(r.itertext())[:40]))

print(f"Tổng số run có màu KHÁC ĐEN (yêu cầu = 0): {non_black_runs}")
if examples_non_black:
    print("Mẫu non-black:", examples_non_black)

# 4. Kiểm tra trích dẫn [1] .. [20]
print("\n--- 4. TRÍCH DẪN VÀ HYPERLINKS ---")
citations_found = 0
for p in doc.paragraphs:
    for hl in p._p.xpath('.//w:hyperlink'):
        anchor = hl.get(qn('w:anchor'))
        if anchor and anchor.startswith('ref_'):
            citations_found += 1

print(f"Tổng số hyperlinks trích dẫn trỏ đến Tài liệu tham khảo: {citations_found}")

print("\nKIỂM TRA HOÀN TẤT!")
