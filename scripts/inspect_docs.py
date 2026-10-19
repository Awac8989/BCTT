import sys
import docx

def main():
    doc = docx.Document('V2_DoamMinhQuan_2224801030137_BCTT (2).docx')
    with open('v2_dump.txt', 'w', encoding='utf-8') as f:
        f.write(f'TOTAL PARAGRAPHS: {len(doc.paragraphs)}\n')
        f.write(f'TOTAL TABLES: {len(doc.tables)}\n\n')
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if t:
                f.write(f'P{i:03d} [{p.style.name}]: {t}\n')
        for t_idx, table in enumerate(doc.tables):
            f.write(f'\n--- TABLE {t_idx} (rows: {len(table.rows)}, cols: {len(table.columns)}) ---\n')
            for r_idx, row in enumerate(table.rows[:10]):
                row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                f.write(f'  Row {r_idx}: {" | ".join(row_txt[:6])}\n')
            if len(table.rows) > 10:
                f.write(f'  ... and {len(table.rows) - 10} more rows\n')

    # Also inspect Mau Bao Cao
    doc_mau = docx.Document('Mau Bao Cao_Thực Tập.docx')
    with open('mau_dump.txt', 'w', encoding='utf-8') as f:
        f.write(f'TOTAL PARAGRAPHS: {len(doc_mau.paragraphs)}\n')
        f.write(f'TOTAL TABLES: {len(doc_mau.tables)}\n\n')
        for i, p in enumerate(doc_mau.paragraphs):
            t = p.text.strip()
            if t:
                f.write(f'P{i:03d} [{p.style.name}]: {t}\n')
        for t_idx, table in enumerate(doc_mau.tables):
            f.write(f'\n--- TABLE {t_idx} (rows: {len(table.rows)}, cols: {len(table.columns)}) ---\n')
            for r_idx, row in enumerate(table.rows[:10]):
                row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                f.write(f'  Row {r_idx}: {" | ".join(row_txt[:6])}\n')

if __name__ == '__main__':
    main()
