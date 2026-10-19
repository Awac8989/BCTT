import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.oxml.ns import qn

for path in ['Mau Bao Cao_Thực Tập.docx', 'V2_DoamMinhQuan_2224801030137_BCTT (2).docx']:
    doc = docx.Document(path)
    print("="*50)
    print("Path:", path)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'v': 'urn:schemas-microsoft-com:vml',
          'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for i in range(min(15, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        drawings = p._p.xpath('.//w:drawing') + p._p.xpath('.//w:pict')
        if drawings:
            print(f"  P[{i}] text='{p.text[:30]}' has drawings: {len(drawings)}")
